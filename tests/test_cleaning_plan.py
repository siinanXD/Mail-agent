"""Putzplan: Wochenmatrix und Excel-Export."""

from __future__ import annotations

import zipfile
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from app.email.importer import import_directory
from app.reports.cleaning_plan import (
    ARRIVAL,
    DEPARTURE,
    OCCUPIED,
    TURNOVER,
    CleaningPlan,
    DayCell,
    UnitRow,
    build_plan,
    export_cleaning_plan,
    week_range,
    write_workbook,
)
from tests.fakes import rule_based_extractor

# KW 37/2026 = Mo 07.09. bis So 13.09.
YEAR, WEEK = 2026, 37


@pytest.fixture
def seeded(session, sample_dir):
    import_directory(
        session, sample_dir, extractor=rule_based_extractor, embedder=lambda c: []
    )
    session.commit()
    return session


def row_of(plan, name: str):
    return next(row for row in plan.rows if row.unit_name == name)


def test_week_range_liefert_montag_bis_sonntag():
    start, end = week_range(YEAR, WEEK)

    assert start == date(2026, 9, 7)
    assert end == date(2026, 9, 13)
    assert start.strftime("%a") == "Mon"


def test_wechseltag_wird_erkannt(seeded):
    """Meier reist am 12.09. ab, Yilmaz zieht am selben Tag ein."""
    plan = build_plan(seeded, year=YEAR, week=WEEK)
    seeblick = row_of(plan, "Ferienwohnung Seeblick")

    assert seeblick.days[0].status == ARRIVAL  # Mo 07.09. Meier
    assert seeblick.days[1].status == OCCUPIED
    assert seeblick.days[5].status == TURNOVER  # Sa 12.09. Wechsel
    assert seeblick.days[6].status == OCCUPIED  # So 13.09. Yilmaz
    assert set(seeblick.days[5].guests) == {"Familie Meier", "Emre Yilmaz"}
    assert seeblick.cleaning_days == [date(2026, 9, 12)]


def test_abreise_und_umgebuchte_anreise(seeded):
    plan = build_plan(seeded, year=YEAR, week=WEEK)

    # Kanonischer Name ist die zuerst gesehene Schreibweise ("FeWo Bergblick",
    # Mail vom 05.08.) - die spaetere "Ferienwohnung Bergblick" landet hier mit.
    bergblick = row_of(plan, "FeWo Bergblick")
    assert bergblick.days[0].status == OCCUPIED  # laufender Aufenthalt
    assert bergblick.days[5].status == DEPARTURE  # Sa 12.09. Kowalski ab

    anna = row_of(plan, "Haus Anna")
    assert anna.days[2].status == ""  # Mi 09.09. - Umbuchung auf Do
    assert anna.days[3].status == ARRIVAL  # Do 10.09. Novak an
    assert anna.days[6].status == DEPARTURE  # So 13.09. Novak ab


def test_stornierte_buchung_taucht_nicht_auf(seeded):
    """Berger (12.-15.09., storniert) darf keine Reinigung ausloesen."""
    plan = build_plan(seeded, year=YEAR, week=WEEK)
    seeblick = row_of(plan, "Ferienwohnung Seeblick")

    assert "Thomas Berger" not in [
        guest for cell in seeblick.days for guest in cell.guests
    ]
    assert plan.cleaning_count == 3


def test_leere_woche_zeigt_alle_objekte_ohne_termine(seeded):
    plan = build_plan(seeded, year=YEAR, week=1)

    assert plan.cleaning_count == 0
    assert {row.unit_name for row in plan.rows} == {
        "Ferienwohnung Seeblick",
        "FeWo Bergblick",
        "Haus Anna",
    }


def test_namen_aus_mails_werden_in_excel_nie_zur_formel(tmp_path):
    """openpyxl machte aus "=..." eine Formel - Objekt- und Gastnamen kommen aus Mails."""
    boese = '=HYPERLINK("http://angreifer.example","Klick")'
    start, end = week_range(YEAR, WEEK)
    tage = [DayCell(start + timedelta(days=offset)) for offset in range(7)]
    tage[0].status = ARRIVAL
    tage[0].guests = ["=1+1"]
    plan = CleaningPlan(year=YEAR, week=WEEK, start=start, end=end, rows=[UnitRow(boese, tage)])

    path = write_workbook(plan, tmp_path / "plan.xlsx")

    with zipfile.ZipFile(path) as archiv:
        blatt = archiv.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "<f>" not in blatt
    sheet = load_workbook(path).active
    assert sheet["A4"].value == boese
    assert sheet["A4"].data_type == "s"


def test_excel_wird_geschrieben_und_ist_lesbar(seeded, tmp_path):
    plan, path = export_cleaning_plan(
        seeded, year=YEAR, week=WEEK, directory=tmp_path
    )

    assert path.exists()
    assert path.name == "putzplan_KW37_2026.xlsx"

    sheet = load_workbook(path).active
    assert sheet.title == "KW 37"
    assert "Putzplan KW 37/2026" in sheet["A1"].value
    assert sheet["A3"].value == "Objekt"
    assert sheet["B3"].value.startswith("Mo 07.09")

    objekte = [sheet.cell(row=r, column=1).value for r in range(4, 4 + len(plan.rows))]
    assert "Ferienwohnung Seeblick" in objekte

    zeile = objekte.index("Ferienwohnung Seeblick") + 4
    assert sheet.cell(row=zeile, column=7).value.startswith(TURNOVER)  # Sa
