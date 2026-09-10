"""Putzplan als Excel: Wochenbelegung je Objekt (Mo-So).

Aufbau der Matrix: eine Zeile je Objekt, eine Spalte je Wochentag. Jede Zelle
zeigt den Status des Tages, farblich unterlegt:

    WECHSEL   Abreise und Anreise am selben Tag -> Reinigung zwingend
    ABREISE   Gast reist ab                     -> Reinigung faellig
    ANREISE   Gast reist an
    belegt    laufender Aufenthalt
    (leer)    frei
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.config import get_settings
from app.database import repositories as repo
from app.tenancy import tenant_id_for

logger = logging.getLogger(__name__)

WEEKDAYS = ("Mo", "Di", "Mi", "Do", "Fr", "Sa", "So")

FREE = ""
OCCUPIED = "belegt"
ARRIVAL = "ANREISE"
DEPARTURE = "ABREISE"
TURNOVER = "WECHSEL"

#: Reinigung faellig - diese Tage sind der eigentliche Zweck des Plans.
CLEANING_STATES = (DEPARTURE, TURNOVER)

_FILLS = {
    TURNOVER: PatternFill("solid", fgColor="FFC7CE"),
    DEPARTURE: PatternFill("solid", fgColor="FFEB9C"),
    ARRIVAL: PatternFill("solid", fgColor="C6EFCE"),
    OCCUPIED: PatternFill("solid", fgColor="DDEBF7"),
}
_HEADER_FILL = PatternFill("solid", fgColor="44546A")
_THIN = Side(style="thin", color="B0B0B0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


@dataclass
class DayCell:
    day: date
    status: str = FREE
    guests: list[str] = field(default_factory=list)

    @property
    def needs_cleaning(self) -> bool:
        return self.status in CLEANING_STATES


@dataclass
class UnitRow:
    unit_name: str
    days: list[DayCell]

    @property
    def cleaning_days(self) -> list[date]:
        return [cell.day for cell in self.days if cell.needs_cleaning]


@dataclass
class CleaningPlan:
    year: int
    week: int
    start: date
    end: date
    rows: list[UnitRow]

    @property
    def cleaning_count(self) -> int:
        return sum(len(row.cleaning_days) for row in self.rows)


def week_range(year: int, week: int) -> tuple[date, date]:
    """Montag und Sonntag einer ISO-Kalenderwoche."""
    monday = date.fromisocalendar(year, week, 1)
    return monday, monday + timedelta(days=6)


def build_plan(session: Session, *, year: int, week: int) -> CleaningPlan:
    """Baut die Wochenmatrix aus den Buchungen der Datenbank."""
    start, end = week_range(year, week)
    bookings = repo.bookings_in_period(session, start=start, end=end)

    days = [start + timedelta(days=offset) for offset in range(7)]
    rows: dict[str, UnitRow] = {}

    for booking in bookings:
        name = booking.unit.name if booking.unit else "Nicht zugeordnet"
        row = rows.get(name)
        if row is None:
            row = UnitRow(name, [DayCell(day) for day in days])
            rows[name] = row

        for cell in row.days:
            _apply_booking(cell, booking)

    # Objekte ohne Buchung in dieser Woche trotzdem zeigen - leere Zeile ist
    # die Information "diese Woche nichts zu tun".
    for unit in repo.list_units(session):
        rows.setdefault(unit.name, UnitRow(unit.name, [DayCell(day) for day in days]))

    ordered = [rows[name] for name in sorted(rows, key=_sort_key)]
    return CleaningPlan(year=year, week=week, start=start, end=end, rows=ordered)


def _sort_key(name: str) -> tuple[int, str]:
    # "Nicht zugeordnet" ans Ende.
    return (1, name) if name == "Nicht zugeordnet" else (0, name.lower())


def _apply_booking(cell: DayCell, booking) -> None:
    arrival, departure = booking.arrival_date, booking.departure_date
    if not (arrival <= cell.day <= departure):
        return

    if cell.day == arrival and cell.day == departure:
        status = TURNOVER
    elif cell.day == departure:
        status = TURNOVER if cell.status == ARRIVAL else DEPARTURE
    elif cell.day == arrival:
        status = TURNOVER if cell.status == DEPARTURE else ARRIVAL
    else:
        status = OCCUPIED if cell.status == FREE else cell.status

    cell.status = status
    if booking.guest_name and booking.guest_name not in cell.guests:
        cell.guests.append(booking.guest_name)


def write_workbook(plan: CleaningPlan, path: Path) -> Path:
    """Schreibt den Plan als .xlsx."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"KW {plan.week}"

    sheet["A1"] = (
        f"Putzplan KW {plan.week}/{plan.year} "
        f"({plan.start.strftime('%d.%m.')} - {plan.end.strftime('%d.%m.%Y')})"
    )
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)

    headers = ["Objekt"] + [
        f"{name} {(plan.start + timedelta(days=index)).strftime('%d.%m.')}"
        for index, name in enumerate(WEEKDAYS)
    ]
    for column, title in enumerate(headers, start=1):
        cell = sheet.cell(row=3, column=column, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = _BORDER

    for row_index, row in enumerate(plan.rows, start=4):
        label = sheet.cell(row=row_index, column=1, value=row.unit_name)
        label.font = Font(bold=True)
        label.border = _BORDER

        for column_index, day_cell in enumerate(row.days, start=2):
            text = day_cell.status
            if day_cell.guests and day_cell.status != FREE:
                text = f"{day_cell.status}\n{', '.join(day_cell.guests)}"
            cell = sheet.cell(row=row_index, column=column_index, value=text)
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = _BORDER
            if day_cell.status in _FILLS:
                cell.fill = _FILLS[day_cell.status]
            if day_cell.needs_cleaning:
                cell.font = Font(bold=True)

    legend_row = len(plan.rows) + 5
    sheet.cell(row=legend_row, column=1, value="Legende:").font = Font(bold=True)
    for offset, (status, description) in enumerate(
        (
            (TURNOVER, "Abreise + Anreise am selben Tag - Reinigung zwingend"),
            (DEPARTURE, "Abreise - Reinigung faellig"),
            (ARRIVAL, "Anreise"),
            (OCCUPIED, "laufender Aufenthalt"),
        ),
        start=1,
    ):
        marker = sheet.cell(row=legend_row + offset, column=1, value=status)
        marker.fill = _FILLS[status]
        marker.border = _BORDER
        sheet.cell(row=legend_row + offset, column=2, value=description)

    sheet.cell(
        row=legend_row + 6,
        column=1,
        value=f"Reinigungen gesamt: {plan.cleaning_count}",
    ).font = Font(bold=True)

    sheet.column_dimensions["A"].width = 26
    for column in range(2, 9):
        sheet.column_dimensions[get_column_letter(column)].width = 18
    for row_index in range(4, len(plan.rows) + 4):
        sheet.row_dimensions[row_index].height = 34
    sheet.freeze_panes = "B4"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    logger.info("Putzplan geschrieben: %s", path)
    return path


def exports_dir_for(tenant_id: int) -> Path:
    """Eigener Export-Ordner je Mandant."""
    return Path(get_settings().exports_dir) / f"tenant-{tenant_id}"


def export_cleaning_plan(
    session: Session, *, year: int, week: int, directory: Path | None = None
) -> tuple[CleaningPlan, Path]:
    """Baut den Plan und schreibt ihn in den Ordner des Mandanten der Session."""
    plan = build_plan(session, year=year, week=week)
    target = Path(directory) if directory else exports_dir_for(tenant_id_for(session))
    path = target / f"putzplan_KW{week:02d}_{year}.xlsx"
    return plan, write_workbook(plan, path)
